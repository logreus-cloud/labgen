"""Преобразование Markdown в документ Word по правилам ГОСТ 7.32-2017."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import yaml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Cm, Mm, Pt
from markdown_it import MarkdownIt
from markdown_it.token import Token

from .config import Config
from .docxutil import (
    ALIGN,
    add_field,
    add_page_break,
    set_font,
    set_paragraph_format,
    setup_page,
    setup_styles,
    text_width_mm,
)
from .titlepage import render_title_page

FRONT_MATTER_RE = re.compile(r"\A---\s*\n(.*?)\n---\s*(?:\n|\Z)", re.DOTALL)

CAPTION_KINDS = {
    "таблица": "table",
    "table": "table",
    "рисунок": "figure",
    "рис": "figure",
    "figure": "figure",
    "листинг": "listing",
    "listing": "listing",
    "код": "listing",
}
CAPTION_RE = re.compile(
    r"^\s*(" + "|".join(CAPTION_KINDS) + r")\s*:\s*(.+?)\s*$", re.IGNORECASE
)
FENCE_CAPTION_RE = re.compile(r"caption\s*=\s*\"([^\"]*)\"|caption\s*=\s*'([^']*)'")

BULLET_MARKER = "–"  # тире для перечислений


def split_front_matter(text: str) -> tuple[dict[str, Any], str]:
    """Отделяет YAML front matter от тела документа."""
    match = FRONT_MATTER_RE.match(text)
    if not match:
        return {}, text
    data = yaml.safe_load(match.group(1)) or {}
    if not isinstance(data, dict):
        return {}, text
    return data, text[match.end():]


def normalize_front_matter(data: dict[str, Any]) -> dict[str, Any]:
    """Ключи верхнего уровня, кроме служебных, считаются полями титульного листа:
    так в самом отчёте достаточно написать `topic: ...`, не вкладывая его в title_page."""
    known = {"format", "title_page", "toc"}
    result = {key: value for key, value in data.items() if key in known}
    extra = {key: value for key, value in data.items() if key not in known}
    if extra:
        result["title_page"] = {**extra, **(result.get("title_page") or {})}
    return result


def plain_text(token: Token | None) -> str:
    """Текст inline-токена без разметки."""
    if token is None:
        return ""
    if not token.children:
        return token.content
    parts: list[str] = []
    for child in token.children:
        if child.type in ("text", "code_inline"):
            parts.append(child.content)
        elif child.type in ("softbreak", "hardbreak"):
            parts.append(" ")
    return "".join(parts)


@dataclass
class RenderResult:
    figures: int = 0
    tables: int = 0
    listings: int = 0
    warnings: list[str] = field(default_factory=list)


class Renderer:
    """Собирает docx-документ из потока токенов markdown-it."""

    def __init__(self, config: Config, base_dir: Path) -> None:
        self.config = config
        self.fmt = config.fmt
        self.base_dir = base_dir
        self.document = Document()
        self.counters = {"figure": 0, "table": 0, "listing": 0}
        self.section_numbers = [0] * 6
        self.pending_caption: tuple[str, str] | None = None
        self.result = RenderResult()
        self._body_started = False

    # --- публичный API -------------------------------------------------

    def build(self, markdown_text: str, *, title_page: bool = True,
              toc: bool | None = None) -> Document:
        toc = self.config.toc if toc is None else toc
        setup_styles(self.document, self.fmt)
        setup_page(self.document, self.fmt, title_page=title_page)

        if title_page:
            render_title_page(self.document, self.config.title_page, self.fmt)
            add_page_break(self.document)
        if toc:
            self._add_toc()
            add_page_break(self.document)

        # linkify не включаем: ссылки в отчёте остаются обычным текстом,
        # зато не появляется лишняя зависимость.
        parser = MarkdownIt("commonmark").enable(["table", "strikethrough"])
        tokens = parser.parse(markdown_text)
        self._walk(tokens, 0, len(tokens))
        return self.document

    # --- обход токенов -------------------------------------------------

    def _walk(self, tokens: list[Token], start: int, end: int, list_level: int = 0) -> None:
        index = start
        while index < end:
            token = tokens[index]
            kind = token.type
            if kind == "heading_open":
                close = self._find_close(tokens, index, "heading_close", end)
                self._heading(token, tokens[index + 1])
                index = close + 1
            elif kind == "paragraph_open":
                close = self._find_close(tokens, index, "paragraph_close", end)
                self._paragraph(tokens[index + 1], list_level=list_level)
                index = close + 1
            elif kind in ("fence", "code_block"):
                self._fence(token)
                index += 1
            elif kind in ("bullet_list_open", "ordered_list_open"):
                close = self._find_close(tokens, index, kind.replace("_open", "_close"), end)
                self._list(tokens, index + 1, close, ordered=kind.startswith("ordered"),
                           level=list_level + 1, start_number=int(token.attrGet("start") or 1))
                index = close + 1
            elif kind == "table_open":
                close = self._find_close(tokens, index, "table_close", end)
                self._table(tokens, index, close)
                index = close + 1
            elif kind == "blockquote_open":
                close = self._find_close(tokens, index, "blockquote_close", end)
                self._walk(tokens, index + 1, close, list_level=list_level + 1)
                index = close + 1
            else:
                index += 1

    @staticmethod
    def _find_close(tokens: list[Token], open_index: int, close_type: str, end: int) -> int:
        level = tokens[open_index].level
        for index in range(open_index + 1, end):
            if tokens[index].type == close_type and tokens[index].level == level:
                return index
        return end - 1

    # --- блоки ---------------------------------------------------------

    def _add_toc(self) -> None:
        heading = self.document.add_paragraph(style="Heading 1")
        run = heading.add_run("Содержание")
        run.bold = True
        set_font(run, self.fmt["font"], self.fmt["size"])
        paragraph = self.document.add_paragraph()
        set_paragraph_format(paragraph, self.fmt, indent=False, alignment="left")
        add_field(
            paragraph,
            'TOC \\o "1-3" \\h \\z \\u',
            placeholder="Нажмите F9 в Word, чтобы собрать содержание.",
        )

    def _is_unnumbered(self, text: str) -> bool:
        normalized = text.strip().lower().rstrip(".")
        return any(normalized.startswith(prefix) for prefix in self.fmt["unnumbered_headings"])

    def _heading(self, opener: Token, inline: Token) -> None:
        level = min(int(opener.tag[1]), 4)
        text = plain_text(inline)
        prefix = ""
        if self.fmt["number_headings"] and not self._is_unnumbered(text):
            self.section_numbers[level - 1] += 1
            for deeper in range(level, len(self.section_numbers)):
                self.section_numbers[deeper] = 0
            prefix = ".".join(str(number) for number in self.section_numbers[:level]) + " "

        paragraph = self.document.add_paragraph(style=f"Heading {level}")
        if level == 1 and self.fmt["section_page_break"] and self._body_started:
            paragraph.paragraph_format.page_break_before = True
        self._body_started = True

        if prefix:
            run = paragraph.add_run(prefix)
            run.bold = True
            set_font(run, self.fmt["font"], self.fmt["size"])
        self._inline(paragraph, inline, bold=True)

    def _paragraph(self, inline: Token, *, list_level: int = 0) -> None:
        caption = self._match_caption(inline)
        if caption:
            self.pending_caption = caption
            return

        image = self._only_image(inline)
        if image is not None:
            self._figure(image)
            return

        self._body_started = True
        paragraph = self.document.add_paragraph()
        set_paragraph_format(paragraph, self.fmt, indent=list_level == 0)
        if list_level:
            paragraph.paragraph_format.left_indent = Cm(self.fmt["first_line_indent"] * list_level)
        self._inline(paragraph, inline)

    def _match_caption(self, inline: Token) -> tuple[str, str] | None:
        match = CAPTION_RE.match(plain_text(inline))
        if not match:
            return None
        return CAPTION_KINDS[match.group(1).lower()], match.group(2)

    @staticmethod
    def _only_image(inline: Token) -> Token | None:
        children = inline.children or []
        images = [child for child in children if child.type == "image"]
        others = [
            child
            for child in children
            if child.type not in ("image", "softbreak") and child.content.strip()
        ]
        return images[0] if len(images) == 1 and not others else None

    def _take_caption(self, kind: str) -> str | None:
        if self.pending_caption and self.pending_caption[0] == kind:
            text = self.pending_caption[1]
            self.pending_caption = None
            return text
        return None

    def _caption_paragraph(self, kind: str, text: str, *, above: bool) -> None:
        self.counters[kind] += 1
        label = self.fmt["captions"][kind]
        dash = self.fmt["captions"]["dash"]
        body = f"{label} {self.counters[kind]}"
        if text:
            body = f"{body} {dash} {text}"
        paragraph = self.document.add_paragraph()
        set_paragraph_format(
            paragraph,
            self.fmt,
            indent=False,
            alignment="left" if above else "center",
            spacing=1.0 if above else None,
        )
        paragraph.paragraph_format.keep_with_next = above
        run = paragraph.add_run(body)
        set_font(run, self.fmt["font"], self.fmt["size"])

    def _figure(self, image: Token) -> None:
        self._body_started = True
        src = image.attrGet("src") or ""
        path = (self.base_dir / src).resolve()
        caption = self._take_caption("figure") or image.content or ""

        paragraph = self.document.add_paragraph()
        set_paragraph_format(paragraph, self.fmt, indent=False, alignment="center", spacing=1.0)
        paragraph.paragraph_format.keep_with_next = True
        if not path.exists():
            self.result.warnings.append(f"изображение не найдено: {src}")
            run = paragraph.add_run(f"[нет изображения: {src}]")
            set_font(run, self.fmt["font"], self.fmt["size"])
        else:
            paragraph.add_run().add_picture(str(path))
            shape = self.document.inline_shapes[-1]
            max_width = Mm(text_width_mm(self.fmt))
            if shape.width > max_width:
                ratio = max_width / shape.width
                shape.width = max_width
                shape.height = int(shape.height * ratio)

        self._caption_paragraph("figure", caption, above=False)

    def _fence(self, token: Token) -> None:
        self._body_started = True
        caption = self._take_caption("listing")
        if caption is None:
            match = FENCE_CAPTION_RE.search(token.info or "")
            if match:
                caption = match.group(1) or match.group(2)
        if caption is not None:
            self._caption_paragraph("listing", caption, above=True)

        code_cfg = self.fmt["code"]
        paragraph = self.document.add_paragraph()
        pf = paragraph.paragraph_format
        pf.alignment = ALIGN["left"]
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(6)
        pf.space_after = Pt(6)
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = code_cfg.get("line_spacing", 1.0)

        lines = token.content.rstrip("\n").split("\n")
        for number, line in enumerate(lines):
            run = paragraph.add_run(line.replace("\t", "    "))
            set_font(run, code_cfg["font"], code_cfg["size"])
            if number < len(lines) - 1:
                run.add_break()

    def _list(self, tokens: list[Token], start: int, end: int, *, ordered: bool,
              level: int, start_number: int) -> None:
        self._body_started = True
        number = start_number
        index = start
        while index < end:
            if tokens[index].type != "list_item_open":
                index += 1
                continue
            close = self._find_close(tokens, index, "list_item_close", end)
            marker = f"{number})" if ordered else BULLET_MARKER
            self._list_item(tokens, index + 1, close, marker=marker, level=level)
            number += 1
            index = close + 1

    def _list_item(self, tokens: list[Token], start: int, end: int, *,
                   marker: str, level: int) -> None:
        index = start
        first = True
        while index < end:
            token = tokens[index]
            if token.type == "paragraph_open":
                close = self._find_close(tokens, index, "paragraph_close", end)
                paragraph = self.document.add_paragraph()
                set_paragraph_format(paragraph, self.fmt, indent=False)
                paragraph.paragraph_format.left_indent = Cm(
                    self.fmt["first_line_indent"] * level
                )
                if first:
                    run = paragraph.add_run(f"{marker} ")
                    set_font(run, self.fmt["font"], self.fmt["size"])
                    first = False
                self._inline(paragraph, tokens[index + 1])
                index = close + 1
            elif token.type in ("bullet_list_open", "ordered_list_open"):
                close = self._find_close(
                    tokens, index, token.type.replace("_open", "_close"), end
                )
                self._list(tokens, index + 1, close, ordered=token.type.startswith("ordered"),
                           level=level + 1, start_number=int(token.attrGet("start") or 1))
                index = close + 1
            elif token.type == "fence":
                self._fence(token)
                index += 1
            else:
                index += 1

    def _table(self, tokens: list[Token], start: int, end: int) -> None:
        self._body_started = True
        rows: list[list[Token | None]] = []
        header_rows = 0
        in_head = False
        index = start
        while index <= end:
            token = tokens[index]
            if token.type == "thead_open":
                in_head = True
            elif token.type == "thead_close":
                in_head = False
            elif token.type == "tr_open":
                close = self._find_close(tokens, index, "tr_close", end + 1)
                row: list[Token | None] = []
                cursor = index + 1
                while cursor < close:
                    if tokens[cursor].type in ("th_open", "td_open"):
                        cell_close = self._find_close(
                            tokens, cursor, tokens[cursor].type.replace("_open", "_close"), close
                        )
                        inline = tokens[cursor + 1] if tokens[cursor + 1].type == "inline" else None
                        row.append(inline)
                        cursor = cell_close + 1
                    else:
                        cursor += 1
                rows.append(row)
                if in_head:
                    header_rows += 1
                index = close
            index += 1

        if not rows:
            return

        caption = self._take_caption("table")
        if caption is not None:
            self._caption_paragraph("table", caption, above=True)

        columns = max(len(row) for row in rows)
        table = self.document.add_table(rows=0, cols=columns)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_size = self.fmt["table"]["size"]

        for row_index, row in enumerate(rows):
            cells = table.add_row().cells
            for column_index in range(columns):
                paragraph = cells[column_index].paragraphs[0]
                set_paragraph_format(
                    paragraph,
                    self.fmt,
                    indent=False,
                    alignment="center" if row_index < header_rows else "left",
                    spacing=1.0,
                )
                inline = row[column_index] if column_index < len(row) else None
                if inline is not None:
                    self._inline(paragraph, inline, bold=row_index < header_rows, size=table_size)

    # --- inline --------------------------------------------------------

    def _inline(self, paragraph, inline: Token, *, bold: bool = False,
                size: float | None = None) -> None:
        font = self.fmt["font"]
        size = size or self.fmt["size"]
        bold_depth = 1 if bold else 0
        italic_depth = 0
        for child in inline.children or []:
            kind = child.type
            if kind == "strong_open":
                bold_depth += 1
            elif kind == "strong_close":
                bold_depth = max(0, bold_depth - 1)
            elif kind == "em_open":
                italic_depth += 1
            elif kind == "em_close":
                italic_depth = max(0, italic_depth - 1)
            elif kind == "code_inline":
                run = paragraph.add_run(child.content)
                set_font(run, self.fmt["code"]["font"], size)
                run.bold = bold_depth > 0
            elif kind == "text":
                if not child.content:
                    continue
                run = paragraph.add_run(child.content)
                set_font(run, font, size)
                run.bold = bold_depth > 0
                run.italic = italic_depth > 0
            elif kind == "softbreak":
                run = paragraph.add_run(" ")
                set_font(run, font, size)
            elif kind == "hardbreak":
                paragraph.add_run().add_break()
            elif kind == "image":
                self.result.warnings.append(
                    f"изображение «{child.content}» внутри абзаца пропущено — "
                    "вынесите его в отдельную строку"
                )


def render_markdown(markdown_text: str, config: Config, base_dir: Path, *,
                    title_page: bool = True,
                    toc: bool | None = None) -> tuple[Document, RenderResult]:
    """Собирает документ; возвращает его вместе со статистикой и предупреждениями."""
    front_matter, body = split_front_matter(markdown_text)
    if front_matter:
        config = config.merged_with(normalize_front_matter(front_matter))
    renderer = Renderer(config, base_dir)
    document = renderer.build(body, title_page=title_page, toc=toc)
    renderer.result.figures = renderer.counters["figure"]
    renderer.result.tables = renderer.counters["table"]
    renderer.result.listings = renderer.counters["listing"]
    return document, renderer.result
