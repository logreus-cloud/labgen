from __future__ import annotations

from pathlib import Path

from docx.shared import Cm, Mm, Pt

from labgen.config import Config
from labgen.render import render_markdown, split_front_matter


def build(markdown: str, workdir: Path, **kwargs):
    kwargs.setdefault("title_page", False)
    kwargs.setdefault("toc", False)
    return render_markdown(markdown, Config(), workdir, **kwargs)


def texts(document) -> list[str]:
    return [p.text for p in document.paragraphs if p.text.strip()]


def test_headings_are_numbered_except_the_standard_ones(workdir):
    document, _ = build(
        "# Введение\n\nтекст\n\n# Теория\n\n## Модель\n\n### Детали\n\n# Заключение\n",
        workdir,
    )
    assert texts(document) == [
        "Введение",
        "текст",
        "1 Теория",
        "1.1 Модель",
        "1.1.1 Детали",
        "Заключение",
    ]


def test_sibling_sections_continue_numbering(workdir):
    document, _ = build("# Первый\n\n## A\n\n## B\n\n# Второй\n\n## C\n", workdir)
    assert texts(document) == ["1 Первый", "1.1 A", "1.2 B", "2 Второй", "2.1 C"]


def test_figure_gets_caption_and_number(workdir):
    document, result = build("![Схема алгоритма](images/scheme.png)\n", workdir)
    assert texts(document) == ["Рисунок 1 — Схема алгоритма"]
    assert result.figures == 1
    assert len(document.inline_shapes) == 1


def test_wide_image_is_scaled_to_text_width(workdir):
    document, _ = build("![Широкая](images/scheme.png)\n", workdir)
    shape = document.inline_shapes[0]
    assert shape.width <= Mm(210 - 30 - 15)


def test_missing_image_warns_but_does_not_crash(workdir):
    document, result = build("![Нет файла](images/absent.png)\n", workdir)
    assert result.warnings
    assert "absent.png" in result.warnings[0]


def test_caption_line_is_consumed_by_the_table(workdir):
    markdown = "Таблица: Замеры\n\n| A | B |\n|---|---|\n| 1 | 2 |\n"
    document, result = build(markdown, workdir)
    assert texts(document) == ["Таблица 1 — Замеры"]
    assert result.tables == 1
    assert [cell.text for cell in document.tables[0].rows[0].cells] == ["A", "B"]


def test_listing_caption_from_fence_info(workdir):
    document, result = build('```python caption="Сортировка"\nx = 1\n```\n', workdir)
    assert texts(document)[0] == "Листинг 1 — Сортировка"
    assert result.listings == 1


def test_fence_without_caption_does_not_consume_a_number(workdir):
    markdown = "```python\nx = 1\n```\n\nЛистинг: Второй\n\n```python\ny = 2\n```\n"
    document, result = build(markdown, workdir)
    assert "Листинг 1 — Второй" in texts(document)
    assert result.listings == 1


def test_counters_are_continuous_across_the_document(workdir):
    markdown = (
        "![Первая](images/scheme.png)\n\n"
        "# Раздел\n\n"
        "![Вторая](images/scheme.png)\n"
    )
    document, result = build(markdown, workdir)
    captions = [text for text in texts(document) if text.startswith("Рисунок")]
    assert captions == ["Рисунок 1 — Первая", "Рисунок 2 — Вторая"]
    assert result.figures == 2


def test_lists_use_gost_markers_and_nesting(workdir):
    markdown = "- первый\n- второй\n    - вложенный\n\n1. раз\n2. два\n"
    document, _ = build(markdown, workdir)
    assert texts(document) == [
        "– первый",
        "– второй",
        "– вложенный",
        "1) раз",
        "2) два",
    ]
    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    assert paragraphs[2].paragraph_format.left_indent > paragraphs[1].paragraph_format.left_indent


def test_inline_formatting_is_preserved(workdir):
    document, _ = build("Обычный **жирный** *курсив* `код`.\n", workdir)
    runs = {run.text: run for run in document.paragraphs[0].runs}
    assert runs["жирный"].bold is True
    assert runs["курсив"].italic is True
    assert runs["код"].font.name == "Courier New"


def test_gost_base_formatting(workdir):
    document, _ = build("Текст отчёта.\n", workdir)
    normal = document.styles["Normal"]
    assert normal.font.name == "Times New Roman"
    assert normal.font.size == Pt(14)
    assert normal.paragraph_format.line_spacing == 1.5
    # отступ хранится в твипах, поэтому сравниваем с допуском округления
    assert abs(normal.paragraph_format.first_line_indent - Cm(1.25)) < Cm(0.01)
    section = document.sections[0]
    assert abs(section.left_margin - Mm(30)) < Mm(0.1)
    assert abs(section.right_margin - Mm(15)) < Mm(0.1)


def test_title_page_and_toc_are_optional(workdir):
    with_extras, _ = render_markdown(
        "# Введение\n", Config({"title_page": {"university": "ВУЗ"}}), workdir,
        title_page=True, toc=True,
    )
    assert "ВУЗ" in texts(with_extras)
    assert "Содержание" in texts(with_extras)

    without, _ = build("# Введение\n", workdir)
    assert texts(without) == ["Введение"]


def test_front_matter_overrides_title_page(workdir):
    markdown = (
        "---\ntopic: Тема из отчёта\n"
        "work_kind: по лабораторной работе № 7\n---\n\n# Введение\n"
    )
    document, _ = render_markdown(markdown, Config(), workdir, title_page=True, toc=False)
    assert "на тему «Тема из отчёта»" in texts(document)
    assert "по лабораторной работе № 7" in texts(document)


def test_split_front_matter_without_front_matter():
    data, body = split_front_matter("# Заголовок\n")
    assert data == {}
    assert body == "# Заголовок\n"


def test_numbering_can_be_disabled(workdir):
    document, _ = render_markdown(
        "# Теория\n", Config({"format": {"number_headings": False}}), workdir,
        title_page=False, toc=False,
    )
    assert texts(document) == ["Теория"]
